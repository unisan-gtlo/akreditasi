"""
Buat file import Excel LAMSPAK Unggul dari lamspak_unggul.json.

    python master_akreditasi/data/build_lamspak_unggul_xlsx.py

Hasil: master_akreditasi/data/LAMSPAK-UNGGUL_import.xlsx (format sama dengan
template Import Excel SIAKRED: sheet SubStandar + ButirDokumen).
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
data = json.loads((HERE / "lamspak_unggul.json").read_text(encoding="utf-8"))

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
REVIEW_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
WRAP = Alignment(vertical="top", wrap_text=True)


def header(ws, cols):
    for i, (name, width) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        ws.column_dimensions[c.column_letter].width = width
    ws.freeze_panes = "A2"


wb = Workbook()

# ---------- Petunjuk ----------
ws = wb.active
ws.title = "Petunjuk"
ws.column_dimensions["A"].width = 110
lines = [
    "IMPORT LAMSPAK UNGGUL (11 STANDAR, 214 BUTIR)",
    "",
    "Sumber: dokumen 'STANDAR 1 SAMPAI 11 (UNGGUL LAMSPAK).docx'.",
    "Syarat: instrumen LAMSPAK-UNGGUL + Standar 1-11 sudah dibuat (python manage.py setup_lamspak_unggul).",
    "",
    "SEBELUM UPLOAD, periksa kolom berlatar KUNING di sheet ButirDokumen:",
    "  kategori : siapa yang mengunggah. UNIVERSITAS = level PT (LPM, BAAK, perpustakaan, keuangan PT),",
    "             FAKULTAS = UPPS, PRODI = program studi. Akun PRODI juga bisa mengunggah butir FAKULTAS.",
    "  wajib    : Y / N",
    "  format   : PDF / DOCX / XLSX / PPTX / GAMBAR / APAPUN",
    "  akses    : INTERNAL / TERBUKA (default saat upload, bisa diubah pengunggah)",
    "",
    "Kolom kode_bersama: butir dengan kode sama (juga di instrumen lain) memakai dokumen yang sama,",
    "  cukup diunggah sekali. UNIV-... = dokumen Universitas, UPPS-... = dokumen fakultas/UPPS,",
    "  PRODI-... = dokumen prodi yang diminta di dua standar.",
    "",
    "Upload: menu Master > Import Excel > instrumen 'LAMSPAK Unggul' > mode UPDATE > Preview > Commit.",
    "Sheet Petunjuk ini diabaikan oleh sistem.",
]
for i, t in enumerate(lines, start=1):
    ws.cell(row=i, column=1, value=t).alignment = WRAP
ws["A1"].font = Font(bold=True, size=13)

# ---------- SubStandar ----------
ws = wb.create_sheet("SubStandar")
header(ws, [("nomor_standar", 15), ("nomor_substandar", 18), ("nomor_parent", 14), ("nama", 40), ("deskripsi", 90)])
for r, std in enumerate(data["standar"], start=2):
    ws.append([std["nomor"], f"{std['nomor']}.1", "", data["substandar_nama"],
               f"Bukti dan dokumen pendukung Standar {std['nama']} (penempatan dalam FED)."])

# ---------- ButirDokumen ----------
ws = wb.create_sheet("ButirDokumen")
cols = [("nomor_substandar", 16), ("kode_butir", 11), ("nama_dokumen", 48), ("kategori", 15),
        ("wajib", 8), ("format", 10), ("ukuran_max", 11), ("akses", 11), ("deskripsi", 60),
        ("panduan_dokumen", 60), ("kode_bersama", 34)]
header(ws, cols)
row = 2
for std in data["standar"]:
    for b in std["butir"]:
        ws.append([f"{std['nomor']}.1", b["kode"], b["nama_dokumen"], b["kategori"],
                   "Y" if b["wajib"] else "N", b["format"], 50, "INTERNAL",
                   b["deskripsi"], b["panduan_dokumen"], b.get("kode_bersama", "")])
        for col in (4, 5, 6, 8):
            ws.cell(row=row, column=col).fill = REVIEW_FILL
        row += 1
last = row - 1

for col, options in (("D", "UNIVERSITAS,BIRO,FAKULTAS,PRODI"), ("E", "Y,N"),
                     ("F", "PDF,DOCX,XLSX,PPTX,GAMBAR,APAPUN"), ("H", "INTERNAL,TERBUKA")):
    dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=False)
    dv.add(f"{col}2:{col}{last}")
    ws.add_data_validation(dv)
ws.auto_filter.ref = f"A1:K{last}"

out = HERE / "LAMSPAK-UNGGUL_import.xlsx"
wb.save(out)
print(f"OK: {out} ({last - 1} butir)")
